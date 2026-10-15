"""Generate editable slides; requires python-pptx. Run from any directory."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

BASE = Path(__file__).resolve().parent
DOC = BASE.parent / 'GPU_MULTI_VM_SCHEDULING_PPT_20261008.md'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
NAVY, INK, MUTED = '132943', '243B53', '61758A'
BLUE, TEAL, BG, LINE = '246BCE', '087E8B', 'F4F7FB', 'DBE4EE'
FONT = 'Noto Sans CJK SC'

def box(slide, x, y, w, h, fill=BG, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    sh.line.fill.background() if line is None else None
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
    return sh

def txt(slide, x, y, w, h, value, size=22, color=INK, bold=False):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.04)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, ln in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln; p.font.name = FONT; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(10)
    return sh

sections = DOC.read_text().split('\n## ')
def slide(title, subtitle, source, note_indices):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
    box(s, 0, 0, 13.333, .15, BLUE)
    txt(s,.55,.4,12.2,.6,title,30,NAVY,True)
    txt(s,.6,1.1,12.1,.7,subtitle,18,MUTED)
    txt(s,.6,6.95,11.6,.36,source,11,MUTED)
    txt(s,12.25,6.94,.5,.36,f'{len(prs.slides):02}',12,BLUE,True)
    s.notes_slide.notes_text_frame.text = '\n\n'.join(sections[i] for i in note_indices)
    return s

def card(s,x,y,w,h,title,body,accent=BLUE,size=20):
    box(s,x,y,w,h)
    txt(s,x+.18,y+.14,w-.36,.55,title,23,accent,True)
    txt(s,x+.18,y+.84,w-.36,h-.94,body,size)

def banner(s,text,y=6.08):
    box(s,.6,y,12.1,.67,NAVY)
    txt(s,.8,y+.1,11.7,.48,text,18,'FFFFFF',True)

s=slide('多 VM 共享 GPU：提交轮转与执行调度',
    'RK3588 · Mali-G610 · Client / Proxy · 当前代码分析：2026-10-08',
    '范围：共享虚拟化路径；静态源码分析 + 历史双 Client 成功日志', [1])
card(s,.6,2,5.9,3.6,'控制点 1｜Proxy 提交入口',
    '调度对象：session 的提交 RPC\n策略：FIFO + session round-robin\n结果：请求进入真实 Panthor',BLUE,23)
card(s,6.8,2,5.9,3.6,'控制点 2｜Panthor / CSF',
    '调度对象：GPU Group / Queue\n策略：优先级 + 资源约束 + 同级轮转\n结果：Group 驻留，固件驱动 GPU',TEAL,23)
banner(s,'核心结论：入口轮转保障提交机会；GPU 时间分配由底层执行调度决定。')

s=slide('架构：多租户工作汇入同一个 GPU 调度域',
    '物理 GPU 由 Proxy VM 持有；Host 不在 Client 之间切换整块 GPU 归属。',
    '源码：panthor-client:1773 → panthor-proxy:2608 / 1770 → panthor_drv:2211', [1,3])
card(s,.6,2,2.5,1.5,'Client VM A','Mesa → panthor-client',BLUE,17)
card(s,.6,3.85,2.5,1.5,'Client VM B','Mesa → panthor-client',BLUE,17)
txt(s,3.12,2.65,.55,2.4,'→\n\n→',28,BLUE)
card(s,3.7,2,3.3,3.35,'Proxy VM','独立 RPC channel\nsession 身份 / 句柄隔离\n提交入口 RR\n真实 Panthor 对象',BLUE,19)
txt(s,7.05,3.17,.55,.7,'→',28,TEAL)
card(s,7.65,2,5.05,3.35,'真实 GPU 执行路径','Queue：job 与 fence 依赖\nGroup：CSG / GPU AS 槽调度\nCSF 固件：消费命令流\nShader cores：执行 shader',TEAL,19)
banner(s,'对象层次：Client VM → session → GPU VM / Group → Queue → job')

s=slide('Proxy：每轮只取一个 session 的一个提交 RPC',
    '每 session 一条 FIFO；全局 runnable 列表只保存有待处理请求的 session。',
    '源码：panthor_proxy_drv.c:654 pick / 686 worker / 707 enqueue / 610 close', [2])
card(s,.6,2,5.65,3.85,'选取算法','① 取 runnable 表头 session\n② 取 session FIFO 表头一个请求\n③ 有剩余请求则 session 回到表尾\n④ worker 调用真实 GROUP_SUBMIT\n⑤ 调用返回后唤醒 handler 并回包',BLUE,21)
card(s,6.55,2,6.15,3.85,'算法示意｜假设请求已同时排队',
    'A：A1  A2  A3        B：B1  B2\n顺序：A1 → B1 → A2 → B2 → A3\n\n轮转单位：一个 RPC，可含多个 jobs\n默认通道同步 RPC，深队列并非必然\n全局 pending 上限 256，满则 -EBUSY',TEAL,20)
banner(s,'worker 串行处理提交调用，但不等待 GPU 完成；多个租户 job 可同时 in-flight。')

s=slide('从提交成功到执行完成：两个不同的返回点',
    '真实 Panthor 将提交转成 job；依赖满足后写入 Queue ringbuffer。',
    '源码：panthor_drv:1228 / 1254 / 2211；panthor_sched:3544 / 3334 / 3764', [3])
steps=[('1  对象翻译','session 内的\nGroup / syncobj\n→ 真实句柄'),('2  接纳 job','收集 fence 依赖\narm + push_job\n→ 提交调用返回'),('3  Queue ready','依赖满足后\nringbuffer 写 CALL\n→ Group / doorbell'),('4  GPU 完成','固件完成序号\nIRQ / sync work\n→ signal fence')]
for i,(t,b) in enumerate(steps):
    card(s,.6+i*3.1,2.15,2.8,3.55,t,b,BLUE if i<2 else TEAL,20)
    if i<3: txt(s,3.43+i*3.1,3.35,.28,.6,'→',19,TEAL)
banner(s,'GROUP_SUBMIT RSP 表示提交调用返回；Client 等待 syncobj 才确认 job 完成或报错。')

s=slide('Panthor：按 Group 选择并配置执行上下文',
    'GPU VM 是 GPU 地址空间；CSG 是固件上下文槽，不能理解为 shader core。',
    '源码：panthor_sched.c:2280 pick / 1320 bind / 1648 program / 2647 apply', [4])
card(s,.6,2,3.8,3.85,'选择候选','优先级从高到低\n跳过故障 / 不可运行 Group\n检查 CSG 槽是否有空间\n检查不同 GPU VM 的 AS 槽\n相同 GPU VM 可复用 AS',BLUE,20)
card(s,4.65,2,3.8,3.85,'建立驻留','激活 GPU VM 地址空间\nGroup 绑定 CSG\n配置 Queue / GPU AS\n设置 endpoint / core mask\n通知固件 START / RESUME',TEAL,20)
card(s,8.7,2,4,3.85,'换入与换出','被换出 Group：SUSPEND\n不可运行 Group：TERMINATE\n固件确认后解绑 CSG / AS\n可运行 Group 回队尾\n重新选组并编程上下文',TEAL,20)
banner(s,'复用现有 Panthor 的 priority + round-robin；实际 shader 工作由固件 / 硬件安排。')

s=slide('两种资源场景：为何不总是逐 VM 轮流运行',
    '是否启动周期轮转取决于 Group 竞争与资源状态，而非 Client VM 的数量本身。',
    '源码：panthor_sched.c:2803 resched target / 2838 tick / 4241 default 10 ms', [4,5])
card(s,.6,2,5.9,3.8,'场景 A｜资源足够','GA、GB 可同时驻留 CSG / AS\nCSF 固件继续调度活跃 Group\n无需因两个 VM 强制周期轮转\n\n同时驻留 ≠ 任意时刻 shader 并行',BLUE,22)
card(s,6.8,2,5.9,3.8,'场景 B｜同级 Group 竞争 CSG',
    'CSG 满且存在同级等待 Group\nSUSPEND → 释放槽 → RESUME\n槽 S：GA → GB → GA（示意）\n\n默认 tick 10 ms，非严格运行配额',TEAL,22)
banner(s,'提交 RR 和 Group 同级轮转都不直接保证每 VM 50% 的 GPU 时间。')

s=slide('默认策略与边界：CPU、优先级、GPU core',
    '将环境控制、入口政策和可选实验分开，避免把它们写成同一种调度机制。',
    '源码：run-vmshm-2client-e2e.sh:369；panthor_proxy_drv.c:1037 / 152 / 197', [4,6])
card(s,.6,2,3.8,3.85,'Host CPU 放置','默认四 CPU 条件下\nbroker / Proxy / 两 Client\n分别放在 CPU 0 / 1 / 2 / 3\n降低 CPU 竞争干扰\n退出恢复原 affinity',BLUE,20)
card(s,4.65,2,3.8,3.85,'共享入口优先级','Client 请求只允许\nLOW / MEDIUM\n高于 MEDIUM 返回 -EACCES\n底层 HIGH / RT 不直接开放\n混合优先级非同等待遇',TEAL,20)
card(s,8.7,2,4,3.85,'可选 core-mask 实验','group_core_partitions 默认 0\n配置 ≥2 时按 session 分片\n修改 Group 的 core mask\n未作为默认设计启用\n不能宣称按 VM 硬件 MIG',TEAL,20)
banner(s,'公平性对象仍是 session / Group；一个 VM 创建多个 session / Group 会影响竞争。')

s=slide('成功证据：双 Client 跑通；公平性还需直接测量',
    '2026-07-11 留存运行：两 Client 都使用 Mali-G610 (Panfrost)，采样计算检查通过。',
    '日志：GPU-SFTP/log/shared/vmshm-2client/vmshm-2client-gles-32m-opencca-pmthor-20260711-151421', [7])
card(s,.6,2,3.8,3.65,'Client 0','29.980 ms\n平均迭代端到端时延\n20 次迭代\nCOMPUTE_CHECK = PASS',BLUE,24)
card(s,4.65,2,3.8,3.65,'Client 1','29.987 ms\n平均迭代端到端时延\n20 次迭代\nCOMPUTE_CHECK = PASS',TEAL,24)
card(s,8.7,2,4,3.65,'证据的范围','count = 8,388,608\n正确性检查 samples = 16\n时延包含整个迭代\n不能当成纯 GPU 时间\n本次未重新运行硬件测试',TEAL,19)
banner(s,'支持双 Client 功能运行；尚未直接证明 RR 顺序、Group 切换和租户 GPU 时间公平。')

s=slide('贡献与路线：先建立可观测性，再实现租户策略',
    '现有工作完成跨 VM 提交整合；GPU 时间 SLA 需要入口和执行层共同约束。',
    '依据：当前 session / worker / Group 代码；下一步均为设计建议，尚未实现', [8,9])
card(s,.6,2,3.8,3.9,'当前自研','Client / Proxy 提交桥接\nchannel 派生租户身份\nsession 对象命名空间\nFIFO + RR 提交入口\npending 上限与关闭唤醒\n双 Client 自动化执行',BLUE,20)
card(s,4.65,2,3.8,3.9,'复用底层能力','DRM Queue scheduler / entity\nfence 依赖与完成同步\nPanthor Group 优先级\nCSG / AS 资源约束\n固件 suspend / resume\nCSF / shader 执行',TEAL,20)
card(s,8.7,2,4,3.9,'下一步开发','关联 VMID → session → job\n测入队 / ready / bind / 完成\n超槽、多 session、长短任务\n按 VM 聚合并限制预算\n联动 Group 层租户策略\n用 p95 / p99 与进展性验证',TEAL,20)
banner(s,'CSG 驻留时间只是代理指标；仅加权 RPC 轮转不能保证 GPU 执行时间公平。')

out = BASE / 'GPU_MULTI_VM_SCHEDULING_20261008.pptx'
prs.save(out)
print(out)
print(f'{len(prs.slides)} editable slides; speaker notes included')
